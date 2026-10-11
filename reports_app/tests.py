from io import BytesIO

from django.test import TestCase
from openpyxl import load_workbook

from products.models import Category, Product, Unit
from sales.views import SAMPLE_SALES


class ReportExportTests(TestCase):
    def setUp(self):
        category = Category.objects.create(name='Fabric')
        unit = Unit.objects.create(name='Meter', abbreviation='m')
        Product.objects.create(
            sku='FAB-001',
            name='Cotton Fabric',
            category=category,
            unit=unit,
            purchase_price='125.50',
            current_stock='20',
            minimum_stock='5',
            reorder_qty='10',
        )

    def test_excel_exports_sales_rows_and_attachment(self):
        response = self.client.get('/reports/export/excel/?type=sales')

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertEqual(
            response['Content-Disposition'],
            'attachment; filename="sales_report.xlsx"',
        )
        worksheet = load_workbook(BytesIO(response.content), read_only=True).active
        self.assertEqual(worksheet['A1'].value, 'Invoice #')
        self.assertEqual(worksheet.max_row, len(SAMPLE_SALES) + 1)
        self.assertEqual(worksheet['A2'].value, SAMPLE_SALES[0]['number'])
        self.assertEqual(worksheet['E2'].number_format, '"₨"#,##0.00;[Red]-"₨"#,##0.00')

    def test_excel_exports_current_inventory(self):
        response = self.client.get('/reports/export/excel/?type=inventory')

        self.assertEqual(response.status_code, 200)
        worksheet = load_workbook(BytesIO(response.content), read_only=True).active
        self.assertEqual(worksheet['A1'].value, 'SKU')
        self.assertEqual(worksheet['A2'].value, 'FAB-001')
        self.assertEqual(worksheet['D2'].value, 20)
        self.assertEqual(worksheet['G2'].value, 2510)

    def test_pdf_exports_as_attachment(self):
        response = self.client.get('/reports/export/pdf/?type=inventory')

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertEqual(
            response['Content-Disposition'],
            'attachment; filename="sales_report.pdf"',
        )
        self.assertTrue(response.content.startswith(b'%PDF-'))

    def test_unknown_report_type_is_rejected(self):
        response = self.client.get('/reports/export/pdf/?type=purchases')
        self.assertEqual(response.status_code, 400)
